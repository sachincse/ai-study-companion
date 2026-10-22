"""Extract text (with page numbers) from uploaded notes: PDF, DOCX, TXT, MD."""
from __future__ import annotations

import io
import logging
from dataclasses import dataclass
from pathlib import Path

from app.services.chunking import PageText

SUPPORTED_EXTENSIONS = {".pdf", ".docx", ".txt", ".md", ".markdown"}
MIME_BY_EXT = {
    ".pdf": "application/pdf",
    ".docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    ".txt": "text/plain",
    ".md": "text/markdown",
    ".markdown": "text/markdown",
}


class UnsupportedFileError(ValueError):
    pass


class EmptyDocumentError(ValueError):
    pass


@dataclass
class ParsedDocument:
    pages: list[PageText]
    page_count: int
    char_count: int
    kind: str


def parse_pdf(data: bytes) -> ParsedDocument:
    from pypdf import PdfReader

    logging.getLogger("pypdf").setLevel(logging.ERROR)  # font-encoding chatter is not actionable

    reader = PdfReader(io.BytesIO(data))
    pages: list[PageText] = []
    total = 0
    for i, page in enumerate(reader.pages, start=1):
        try:
            text = page.extract_text() or ""
        except Exception:  # pragma: no cover - pypdf edge cases
            text = ""
        total += len(text.strip())
        pages.append(PageText(page_no=i, text=text))
    if total < 40:
        raise EmptyDocumentError(
            "No extractable text found. This looks like a scanned PDF; run it through OCR first."
        )
    return ParsedDocument(pages=pages, page_count=len(reader.pages), char_count=total, kind="pdf")


def parse_docx(data: bytes) -> ParsedDocument:
    import docx

    document = docx.Document(io.BytesIO(data))
    lines: list[str] = []
    for para in document.paragraphs:
        text = para.text.strip()
        if not text:
            continue
        style = (para.style.name or "").lower() if para.style is not None else ""
        if style.startswith("heading"):
            lines.append(f"\n# {text}\n")
        else:
            lines.append(text)
    for table in document.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells if c.text.strip()]
            if cells:
                lines.append(" | ".join(cells))
    text = "\n".join(lines)
    if len(text.strip()) < 20:
        raise EmptyDocumentError("The DOCX file has no readable text.")
    return ParsedDocument(pages=[PageText(page_no=None, text=text)], page_count=1, char_count=len(text), kind="docx")


def parse_text(data: bytes) -> ParsedDocument:
    text = data.decode("utf-8", errors="replace")
    if len(text.strip()) < 20:
        raise EmptyDocumentError("The text file is empty.")
    return ParsedDocument(pages=[PageText(page_no=None, text=text)], page_count=1, char_count=len(text), kind="text")


def parse_bytes(filename: str, data: bytes) -> ParsedDocument:
    ext = Path(filename).suffix.lower()
    if ext not in SUPPORTED_EXTENSIONS:
        raise UnsupportedFileError(f"Unsupported file type '{ext or 'unknown'}'. Upload PDF, DOCX, TXT or MD.")
    if ext == ".pdf":
        return parse_pdf(data)
    if ext == ".docx":
        return parse_docx(data)
    return parse_text(data)


def parse_file(path: Path) -> ParsedDocument:
    return parse_bytes(path.name, path.read_bytes())
